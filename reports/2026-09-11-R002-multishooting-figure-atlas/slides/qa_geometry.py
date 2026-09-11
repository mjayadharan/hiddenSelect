"""qa_geometry.py — geometric QA of the deck without LibreOffice: flags overlapping shapes,
shapes off-slide, and text boxes whose estimated wrapped height exceeds their box (font-size
aware, Calibri/Georgia average glyph widths), plus the minimum font size used."""
import sys, pathlib
from pptx import Presentation
from pptx.util import Emu, Pt
p = Presentation(sys.argv[1] if len(sys.argv) > 1 else str(pathlib.Path(__file__).with_name("multishooting_story.pptx")))
W, H = p.slide_width, p.slide_height
AVG = {"Georgia": 0.55, "Calibri": 0.47, "Cambria": 0.5}
minfont = 99; issues = 0
for si, s in enumerate(p.slides, 1):
    boxes = []
    for sh in s.shapes:
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        if x < 0 or y < 0 or x + w > W + 10 or y + h > H + 10: print(f"slide {si}: {sh.shape_type} '{sh.name}' off-slide"); issues += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame; est = Emu(0); usable = w - Emu(91440*0.1)
            for para in tf.paragraphs:
                runs = para.runs
                if not runs: est += Pt(8); continue
                size = max((r.font.size.pt if r.font.size else 16) for r in runs); font = runs[0].font.name or "Calibri"
                chars = sum(len(r.text) for r in runs); cpl = max(1, int(usable / Pt(size * AVG.get(font, 0.5))))
                lines = max(1, -(-chars // cpl)); ls = para.line_spacing or 1.0
                est += Pt(size * 1.2 * ls * lines) + (para.space_after or Pt(0))
                minfont = min(minfont, min((r.font.size.pt if r.font.size else 16) for r in runs))
            if est > h * 1.05:
                print(f"slide {si}: text box '{tf.text[:40]}…' estimated height {est/914400:.2f} in > box {h/914400:.2f} in"); issues += 1
            boxes.append((sh, x, y, w, min(h, est)))
        else:
            boxes.append((sh, x, y, w, h))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            if a[0].shape_type == 1 or b[0].shape_type == 1:   # autoshape (side band) — allowed to touch
                continue
            ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]); oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
            if ox > Emu(91440*0.15) and oy > Emu(91440*0.15):
                print(f"slide {si}: overlap {ox/914400:.2f}x{oy/914400:.2f} in between '{a[0].name}' and '{b[0].name}'"); issues += 1
print(f"min font size {minfont} pt; issues {issues}")
