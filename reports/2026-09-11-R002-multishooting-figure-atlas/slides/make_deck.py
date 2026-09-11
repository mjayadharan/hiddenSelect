"""make_deck.py — 8-slide story deck (python-pptx, 16:9, body text >= 14 pt).
Images are the report's own figures (figures/*.png); every number is quoted from the report tables.
Output: slides/multishooting_story.pptx
"""
import pathlib, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = pathlib.Path(__file__).resolve().parent; ROOT = HERE.parent; FIG = ROOT / "figures"
NUM = json.loads((ROOT / "tables" / "table_numbers.json").read_text())
NAVY, INK, PAPER, CREAM = RGBColor(0x1E, 0x27, 0x61), RGBColor(0x22, 0x22, 0x22), RGBColor(0xFA, 0xFA, 0xF7), RGBColor(0xEE, 0xEC, 0xE4)
BLUE, VERM, GREEN, MUTED, WHITE, ICE = RGBColor(0x00, 0x72, 0xB2), RGBColor(0xD5, 0x5E, 0x00), RGBColor(0x00, 0x9E, 0x73), RGBColor(0x66, 0x66, 0x66), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xCA, 0xDC, 0xFC)
HFONT, BFONT = "Georgia", "Calibri"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation(); prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]

def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color

def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h); s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background(); s.shadow.inherit = False; return s

def text(slide, x, y, w, h, runs, size=16, color=INK, font=BFONT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {bold,color,size,italic}) runs."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(6)
        for r in (para if isinstance(para, list) else [(para, {})]):
            t, st = (r if isinstance(r, tuple) else (r, {}))
            run = p.add_run(); run.text = t; f = run.font; f.name = st.get("font", font); f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", False)
            f.color.rgb = st.get("color", color)
    return tb

def bullets(slide, x, y, w, h, items, size=16, color=INK, bullet="▸", gap=8):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = Inches(0.03)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap); p.line_spacing = 1.12
        head, body = it if isinstance(it, tuple) else (None, it)
        r = p.add_run(); r.text = bullet + "  "; r.font.size = Pt(size); r.font.color.rgb = VERM; r.font.name = BFONT
        if head:
            r = p.add_run(); r.text = head + " "; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = BFONT
        r = p.add_run(); r.text = body; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = BFONT
    return tb

def picture(slide, path, x, y, w=None, h=None):
    im = Image.open(path); ar = im.width / im.height
    if w is not None and h is not None:
        if w / h > ar: w = int(h * ar)
        else: h = int(w / ar)
    elif w is None: w = int(h * ar)
    elif h is None: h = int(w / ar)
    return slide.shapes.add_picture(str(path), x, y, w, h)

def stat(slide, x, y, w, big, label, color=NAVY, size=34):
    text(slide, x, y, w, Inches(0.75), big, size=size, color=color, font=HFONT, bold=True)
    text(slide, x, y + Inches(0.72), w, Inches(1.0), label, size=14, color=MUTED)

def content_slide(n, title, kicker=None):
    s = prs.slides.add_slide(blank); bg(s, PAPER)
    rect(s, 0, 0, Inches(0.32), H, NAVY)
    text(s, Inches(0.02), H - Inches(0.7), Inches(0.3), Inches(0.5), str(n), size=14, color=ICE, font=HFONT, align=PP_ALIGN.CENTER)
    text(s, Inches(0.7), Inches(0.35), Inches(12.2), Inches(0.75), title, size=28, color=NAVY, font=HFONT, bold=True)
    if kicker: text(s, Inches(0.7), Inches(1.1), Inches(12.2), Inches(0.75), kicker, size=15, color=MUTED)
    return s

def caption(slide, x, y, w, t, h=0.75):
    text(slide, x, y, w, Inches(h), t, size=14, color=MUTED)
def below(pic, gap=0.08): return pic.top + pic.height + Inches(gap)

# ---------------------------------------------------------------- slide 1: title
s = prs.slides.add_slide(blank); bg(s, NAVY)
rect(s, 0, 0, Inches(0.32), H, VERM)
text(s, Inches(0.9), Inches(0.9), Inches(7.4), Inches(2.6), "Discovering sparse dynamics with multiple shooting and guess propagation", size=40, color=WHITE, font=HFONT, bold=True, spacing=1.05)
text(s, Inches(0.9), Inches(3.6), Inches(7.2), Inches(1.4), "From noisy trajectory samples to the few governing terms — by solving a sequence of easier problems and carrying the answer forward.", size=18, color=ICE)
text(s, Inches(0.9), Inches(5.6), Inches(7.2), Inches(1.2), [[("Manu Jayadharan, Alasdair Hastewell, Niall Mangan", {"bold": True, "color": WHITE, "size": 16})], [("Figures and numbers: report R002 (11 September 2026), FitzHugh–Nagumo, Lotka–Volterra and Lorenz test systems", {"color": ICE, "size": 14})]])
pic = picture(s, FIG / "fig11_landscape_2d_v_v3.png", Inches(8.4), Inches(1.4), w=Inches(4.6))
text(s, Inches(8.4), below(pic), Inches(4.6), Inches(0.8), "Cost landscape vs window size: smooth for short windows, rugged for single shooting", size=14, color=ICE)

# ---------------------------------------------------------------- slide 2: problem
s = content_slide(2, "Noisy samples in, a sparse equation out", "Model discovery: choose which library terms are non-zero and estimate them by matching simulated trajectories to the data")
bullets(s, Inches(0.7), Inches(1.95), Inches(5.6), Inches(4.9), [
    ("Data.", "N = 101 noisy samples of v(t), w(t) (FitzHugh–Nagumo), spacing Δt = 1, 5 % noise per component."),
    ("Ansatz.", "ẋᵢ = Σⱼ pᵢⱼ Θⱼ(x) over all monomials up to degree 3: 20 unknown coefficients, 7 truly non-zero."),
    ("Single shooting.", "Integrate from y₀ over the whole record and fit p. Errors in p and y₀ are amplified along the orbit: the cost is rugged, has a blow-up plateau, and a small wrong p costs 260× the true one."),
    ("Consequence.", "From a random start, the optimiser never finds the model (control arm: error stays at the seed, 1.65)."),
], size=16)
pic = picture(s, FIG / "fig04_multiple_shooting.png", Inches(6.6), Inches(1.95), w=Inches(6.3), h=Inches(4.4))
caption(s, Inches(6.6), below(pic), Inches(6.3), "Same wrong p: single shooting (top) vs restarting every 10 samples (middle); residuals (bottom)")

# ---------------------------------------------------------------- slide 3: theory & gap
s = content_slide(3, "Multiple shooting, and the gap we fill", "Partition [t₀, t_N] into K windows; each window is integrated from the datum at its left node, with no continuity constraint")
text(s, Inches(0.7), Inches(1.95), Inches(6.2), Inches(1.9), [[("Cost with K windows (κ data intervals each)", {"bold": True, "color": NAVY, "size": 16})],
     [("J_K(p) = Σₖ Σ_{tᵢ ∈ [τₖ₋₁, τₖ]} ‖ φ(tᵢ − τₖ₋₁; p, y_{τₖ₋₁}) − yᵢ ‖²", {"size": 17, "font": "Cambria"})],
     [("K = 1 is single shooting; κ = 1 (a node at every datum) is nearly convex.", {"size": 15, "color": MUTED})]])
bullets(s, Inches(0.7), Inches(3.9), Inches(6.2), Inches(3.2), [
    ("Known.", "Textbook multiple shooting makes the node states extra variables and enforces continuity — many more unknowns, constrained optimisation."),
    ("Gap.", "Nobody exploits the smooth-to-rugged family J₁, J₂, …, J_N directly: solve the easy problem first, then remove nodes gradually while carrying the minimiser forward."),
    ("Needed.", "A bound on how much the cost — and its minimiser — moves when nodes are removed, so the sweep stays inside the right basin."),
], size=16)
pic = picture(s, FIG / "fig03_partition.png", Inches(7.2), Inches(1.95), w=Inches(5.8))
caption(s, Inches(7.2), below(pic), Inches(5.8), "Fine partition (top) and the coarser one after deleting every other node (bottom): the elementary step of the sweep")
text(s, Inches(7.2), below(pic, 0.95), Inches(5.8), Inches(1.4), [[("Node removal changes the cost only on the windows next to the removed nodes ", {"bold": True, "color": NAVY}), ("— by an amount that grows with the number of nodes removed and with the extra span ΔT₂ the surviving segment must cover (Proposition 1, next slides).", {})]], size=15)

# ---------------------------------------------------------------- slide 4: method
s = content_slide(4, "Our method: guess propagation across window sizes", "Solve the smooth problem first, then carry each minimiser into the next, more rugged landscape")
pic = picture(s, FIG / "fig06_guess_propagation.png", Inches(0.7), Inches(1.95), w=Inches(7.6))
text(s, Inches(8.6), Inches(1.95), Inches(4.4), Inches(0.5), "Algorithm", size=18, color=NAVY, font=HFONT, bold=True)
bullets(s, Inches(8.6), Inches(2.45), Inches(4.4), Inches(3.0), [
    "p_guess ← random start",
    "for κ = 1, 2, 3, … , N−1 (few nodes removed per step): p^(κ) ← argmin J_κ(p) from p_guess; p_guess ← p^(κ)",
    "polish with single shooting (κ = N−1) from p_guess",
    "sparsity: smooth-ℓ1 penalty γ Σ|pⱼ| inside every J_κ",
], size=15, bullet="→", gap=6)
text(s, Inches(8.6), Inches(5.3), Inches(4.4), Inches(1.9), [[("Why it works: ", {"bold": True, "color": NAVY}), ("the number of local minima along a random line grows from 1.3 (κ = 1) to 7.9 (κ = 100), but consecutive landscapes are close, so the carried guess stays in the right basin.", {})]], size=15)

# ---------------------------------------------------------------- slide 5: theoretical results
s = content_slide(5, "Theoretical results — and how tight they are", "Lipschitz constant L of the vector field, noise η, window length ΔT; strong convexity with modulus μ near the minimiser")
text(s, Inches(0.7), Inches(1.95), Inches(6.0), Inches(5.0), [
    [("Lemma 1 (flow sensitivity).", {"bold": True, "color": NAVY, "size": 16})],
    [("e^{−Lt} ‖x₁−x₂‖ ≤ ‖φ(t;x₁) − φ(t;x₂)‖ ≤ e^{Lt} ‖x₁−x₂‖", {"size": 16, "font": "Cambria"})],
    [("so the noise at a node is amplified by at most e^{LΔT} within a window instead of e^{L t_N} over the record.", {"size": 15, "color": MUTED})],
    [("Parameter error.", {"bold": True, "color": NAVY, "size": 16})],
    [("‖p* − p^(K)‖ ≤ √( (2/μ) [ N ‖η‖²_max (1 + e^{LΔT})² − J_K(p^(K)) ] )", {"size": 16, "font": "Cambria"})],
    [("Proposition 1 (node removal).", {"bold": True, "color": NAVY, "size": 16})],
    [("|Ĵ_K(p) − J_K(p)| ≤ 2 |I_R| C_max ( e^{L(ΔT₁+ΔT₂)} ‖η‖_max + e^{LΔT₁} [ (L̃/L)(e^{LΔT₂} − 1) ‖p − p*‖ + ‖η‖_max ] )", {"size": 15, "font": "Cambria"})],
    [("Design rule: remove few nodes at a time, keep ΔT₂ small, prefer low-noise nodes.", {"size": 15, "color": MUTED})],
], size=16)
pic = picture(s, FIG / "fig08_lemma1_flow_sensitivity.png", Inches(7.0), Inches(1.95), w=Inches(6.0))
caption(s, Inches(7.0), below(pic), Inches(6.0), "Lemma 1 checked on FHN: 24 perturbed orbits (blue) stay below e^{Lt} (vermilion); the bound holds but is loose", h=0.8)
stat(s, Inches(7.0), Inches(5.3), Inches(2.9), "≤ 8×", "measured amplification of a state perturbation on the FHN orbit")
stat(s, Inches(10.1), Inches(5.3), Inches(2.9), "10¹³", "what e^{Lt} with L = 3.05 allows at t = 10: the bounds order the arms, not the sizes")

# ---------------------------------------------------------------- slide 6: numerical results
t1 = NUM
s = content_slide(6, "Numerical results: FitzHugh–Nagumo, 8 starts", "Same optimiser (Nelder–Mead, 2500 iterations per stage), same seeds, same 16-stage schedule, with and without propagation")
pic = picture(s, FIG / "fig15_sweep_headline.png", Inches(0.7), Inches(1.95), w=Inches(8.4))
caption(s, Inches(0.7), below(pic), Inches(8.4), "Cost, parameter error and recovery score along the sweep; thin lines = seeds, thick = median. Blue: guess propagation; vermilion: restart from the seed every stage", h=0.8)
stat(s, Inches(9.4), Inches(1.9), Inches(3.7), f"{t1['T1_final_perr_GP']:.2f} vs {t1['T1_final_perr_control']:.2f}", "final ‖p − p*‖, propagation vs control (seed medians; start ≈ 1.6)")
stat(s, Inches(9.4), Inches(3.55), Inches(3.7), "7 of 8", "seeds where propagation is strictly better at every κ ≥ 2 (the 8th starts on the blow-up plateau)")
stat(s, Inches(9.4), Inches(5.2), Inches(3.7), f"{t1['T1_final_paired_diff']:.2f}", f"paired difference at κ = 100, 95 % bootstrap CI [{t1['T1_final_ci'][0]:.2f}, {t1['T1_final_ci'][1]:.2f}]")
text(s, Inches(0.7), below(pic, 0.95), Inches(8.4), Inches(1.3), [[("Also: ", {"bold": True, "color": NAVY}), ("the parameter error is smallest at κ = 5 (0.22) and drifts to 0.32 at single shooting while the cost keeps falling — the sparsity term and noise bias the final polish, so stop on validation error, not on cost.", {})]], size=15)

# ---------------------------------------------------------------- slide 7: robustness & generality
s = content_slide(7, "What matters, and does it generalise?", "Every variation re-run with the same seeds; numbers are seed medians at the end of the sweep")
x0 = Inches(0.7); w3 = Inches(4.0); gap = Inches(0.2)
for i, (fig, big, lab) in enumerate((("fig21_schedule.png", "0.30 vs 0.76", "error with a dense schedule (44 steps) vs a 1→100 jump: remove few nodes per step"),
                                     ("fig25_optimizers.png", "0.21 / 0.95", "error / recovery score for BFGS with the exact AD gradient + propagation; gradients die on the flat plateau without it"),
                                     ("fig36_lorenz.png", "0.83", "recovery score on Lorenz (30 coefficients, chaotic) with BFGS + propagation; the control blows up in 4 of 4 seeds"))):
    x = x0 + i * (w3 + gap)
    pic = picture(s, FIG / fig, x, Inches(1.95), w=w3, h=Inches(1.75))
    text(s, x, Inches(3.8), w3, Inches(0.7), big, size=28, color=NAVY, font=HFONT, bold=True)
    text(s, x, Inches(4.45), w3, Inches(1.1), lab, size=14, color=MUTED)
bullets(s, Inches(0.7), Inches(5.6), Inches(12.2), Inches(1.6), [
    ("Robust to noise up to 10 %", "(error 0.22–0.36); fails at 20 %. The control fails even at zero noise — the difficulty is the landscape, not the noise."),
    ("Lotka–Volterra:", "error 0.29, score 0.83 with propagation; 5 of 6 control seeds blow up. A graded blow-up penalty rescues starts that begin on the plateau."),
], size=15)

# ---------------------------------------------------------------- slide 8: takeaways
s = prs.slides.add_slide(blank); bg(s, NAVY); rect(s, 0, 0, Inches(0.32), H, VERM)
text(s, Inches(0.9), Inches(0.5), Inches(12), Inches(0.9), "Takeaways and what comes next", size=34, color=WHITE, font=HFONT, bold=True)
tb = bullets(s, Inches(0.9), Inches(1.6), Inches(6.4), Inches(5.3), [
    ("Guess propagation works.", "Sweeping the window size and carrying the minimiser turns an unsolvable single-shooting fit into a reliable one (7 of 8 starts; paired CI excludes zero)."),
    ("Three design rules.", "Remove few nodes per step; use a graded blow-up penalty; use a gradient method (BFGS through the integrator) once the sweep has started."),
    ("Theory is qualitative.", "The Lipschitz bounds order the arms correctly but are loose by orders of magnitude on an attracting orbit; strong convexity is not verified at Nelder–Mead's iterates."),
    ("Code audit.", "Six defects found and fixed (dropped w-residual, truncated Jacobian, wrong gradient, mixed monomial order, ignored data argument, disabled propagation) — every published number re-derived."),
], size=15, color=ICE, gap=10)
for p in tb.text_frame.paragraphs:
    for r in p.runs:
        if r.font.bold: r.font.color.rgb = WHITE
text(s, Inches(7.7), Inches(1.6), Inches(5.0), Inches(0.5), "Next", size=20, color=WHITE, font=HFONT, bold=True)
bullets(s, Inches(7.7), Inches(2.15), Inches(5.0), Inches(3.5), [
    "Node removal chosen by node noise (the Note in Prop. 1), not by a fixed stride",
    "Tighter bounds via the logarithmic norm (μ = 1.17 vs L = 3.05 for FHN)",
    "Validation-based stopping of the sweep instead of running to single shooting",
    "Real data: unknown noise, partial observations",
], size=15, color=ICE, bullet="→", gap=8)
text(s, Inches(7.7), Inches(5.9), Inches(5.0), Inches(1.0), "Report R002: 37 figures, 6 animations, 18/18 gates — reports/2026-09-11-R002-multishooting-figure-atlas/", size=14, color=ICE)

out = HERE / "multishooting_story.pptx"; prs.save(out); print("saved", out)
